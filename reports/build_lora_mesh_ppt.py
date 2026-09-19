#!/usr/bin/env python3
"""Build a presentation for the LoRa mesh simulator project."""

from __future__ import annotations

import csv
import math
import statistics
from pathlib import Path
from typing import Iterable
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    "/Users/bh4me_macair/Documents/深圳大学/电子与信息工程学院logo 0318/电子与信息工程学院ppt模板0318.pptx"
)
OUT_DIR = ROOT / "output"
SCRATCH = ROOT / "scratch" / "ppt_lora_mesh"
PPTX_OUT = OUT_DIR / "lora_mesh_路由仿真汇报_目录仿真版.pptx"
PREVIEW_DIR = OUT_DIR / "lora_mesh_preview_目录仿真版"

W, H = 13.333333, 7.5
EMU_W, EMU_H = 12192000, 6858000

BLUE = RGBColor(20, 104, 184)
DEEP = RGBColor(17, 38, 74)
GREEN = RGBColor(111, 186, 45)
RED = RGBColor(194, 48, 58)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(95, 111, 132)
PALE = RGBColor(238, 246, 255)
WHITE = RGBColor(255, 255, 255)

FONT = "Microsoft YaHei"
FONT_BOLD = "Microsoft YaHei"
PIL_FONT = "/System/Library/Fonts/Hiragino Sans GB.ttc"


def mean_by_protocol(csv_path: Path) -> dict[str, dict[str, float]]:
    rows: dict[str, list[dict[str, str]]] = {}
    with csv_path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows.setdefault(row["protocol"], []).append(row)
    result: dict[str, dict[str, float]] = {}
    for protocol, values in rows.items():
        result[protocol] = {}
        for key in values[0]:
            if key == "protocol":
                continue
            try:
                result[protocol][key] = statistics.mean(float(v[key]) for v in values)
            except ValueError:
                pass
    return result


def summary_by_protocol(csv_path: Path) -> dict[str, dict[str, dict[str, float]]]:
    rows: dict[str, list[dict[str, str]]] = {}
    with csv_path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows.setdefault(row["protocol"], []).append(row)
    result: dict[str, dict[str, dict[str, float]]] = {}
    for protocol, values in rows.items():
        result[protocol] = {}
        for key in values[0]:
            if key == "protocol":
                continue
            try:
                nums = [float(v[key]) for v in values]
            except ValueError:
                continue
            result[protocol][key] = {
                "mean": statistics.mean(nums),
                "sd": statistics.stdev(nums) if len(nums) >= 2 else 0.0,
            }
    return result


def pct_drop(before: float, after: float) -> float:
    return (before - after) / before if before else 0.0


def prepare_assets() -> dict[str, Path]:
    media = SCRATCH / "media"
    media.mkdir(parents=True, exist_ok=True)
    with ZipFile(TEMPLATE) as z:
        for name in z.namelist():
            if name.startswith("ppt/media/") and name.split("/")[-1]:
                target = media / Path(name).name
                target.write_bytes(z.read(name))
    return {
        "cover": media / "image23.jpeg",
        "circuit": media / "image21.jpeg",
        "light_circuit": media / "image29.png",
        "campus": media / "image27.jpeg",
        "logo": media / "image33.png",
        "mark": media / "image13.png",
        "chip": media / "image4.png",
        "check": media / "image7.png",
    }


def clear_slides(prs: Presentation) -> None:
    xml_slides = prs.slides._sldIdLst  # noqa: SLF001
    rel = prs.part.drop_rel
    for slide_id in list(xml_slides):
        rel(slide_id.rId)
        xml_slides.remove(slide_id)


def set_text(shape, text: str, size: int, color=INK, bold=False, align=None, line_spacing=1.05):
    shape.text = text
    shape.text_frame.word_wrap = True
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for p in shape.text_frame.paragraphs:
        if align is not None:
            p.alignment = align
        p.line_spacing = line_spacing
        for r in p.runs:
            r.font.name = FONT_BOLD if bold else FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def add_text(slide, x, y, w, h, text, size, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box, text, size, color, bold, align)
    return box


def add_title(slide, title: str, subtitle: str | None = None, section=""):
    add_text(slide, 0.78, 0.43, 10.5, 0.48, title, 25, DEEP, True)
    if subtitle:
        add_text(slide, 0.8, 0.95, 10.6, 0.32, subtitle, 10, MUTED)
    if section:
        add_text(slide, 11.25, 0.48, 1.1, 0.35, section, 15, BLUE, True, PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.78), Inches(1.25), Inches(0.8), Inches(0.035))
    rule.fill.solid()
    rule.fill.fore_color.rgb = GREEN
    rule.line.fill.background()


def add_logo(slide, assets):
    slide.shapes.add_picture(str(assets["logo"]), Inches(9.55), Inches(6.95), width=Inches(2.6))


def add_footer(slide, assets, n: int):
    add_text(slide, 0.78, 6.98, 3.5, 0.2, "深圳大学电子与信息工程学院", 7, MUTED)
    add_text(slide, 11.95, 6.96, 0.5, 0.22, f"{n:02d}", 8, BLUE, True, PP_ALIGN.RIGHT)


def add_soft_bg(slide, assets):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(249, 252, 255)
    slide.shapes.add_picture(str(assets["light_circuit"]), Inches(7.75), Inches(0.25), width=Inches(5.2))


def bullet_list(slide, x, y, w, items: Iterable[str], size=15, color=INK, gap=0.46):
    for i, item in enumerate(items):
        cy = y + i * gap
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.09), Inches(0.09), Inches(0.09))
        dot.fill.solid()
        dot.fill.fore_color.rgb = GREEN if i % 2 == 0 else BLUE
        dot.line.fill.background()
        add_text(slide, x + 0.22, cy, w, 0.33, item, size, color)


def add_metric(slide, x, y, number, label, color=BLUE):
    add_text(slide, x, y, 1.65, 0.55, number, 27, color, True, PP_ALIGN.CENTER)
    add_text(slide, x - 0.1, y + 0.58, 1.85, 0.45, label, 9, MUTED, False, PP_ALIGN.CENTER)


def add_bar_chart(slide, x, y, w, h, title, categories, s1_name, s1, s2_name, s2, as_percent=False):
    data = CategoryChartData()
    data.categories = categories
    data.add_series(s1_name, s1)
    data.add_series(s2_name, s2)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(8)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(224, 232, 242)
    if as_percent:
        chart.value_axis.tick_labels.number_format = "0%"
    for idx, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = BLUE if idx == 0 else GREEN
        series.has_data_labels = True
        series.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
        series.data_labels.font.size = Pt(7)
        series.data_labels.number_format = "0%" if as_percent else "0.0"
    return chart


def add_path_diagram(slide):
    nodes = [(1.4, 3.55), (2.8, 2.55), (4.35, 3.25), (5.65, 2.25), (7.1, 3.35), (8.35, 2.65)]
    for a, b in zip(nodes, nodes[1:]):
        line = slide.shapes.add_connector(1, Inches(a[0] + 0.19), Inches(a[1] + 0.19), Inches(b[0] + 0.19), Inches(b[1] + 0.19))
        line.line.color.rgb = RGBColor(126, 156, 191)
        line.line.width = Pt(2)
    for i, (x, y) in enumerate(nodes):
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.38), Inches(0.38))
        c.fill.solid()
        c.fill.fore_color.rgb = BLUE if i in (0, len(nodes) - 1) else GREEN
        c.line.color.rgb = WHITE
        c.line.width = Pt(1.5)
    add_text(slide, 1.0, 4.08, 8.2, 0.35, "多跳 mesh：用节点转发补足网关覆盖、地形遮挡与 P2P 通信需求", 14, DEEP, True, PP_ALIGN.CENTER)


def style_shape(shape, fill, line=None, radius=False):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()


def add_callout(slide, x, y, w, h, head, body, color=BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    style_shape(shp, RGBColor(244, 249, 255), RGBColor(214, 228, 244))
    add_text(slide, x + 0.18, y + 0.13, w - 0.36, 0.28, head, 13, color, True)
    add_text(slide, x + 0.18, y + 0.5, w - 0.36, h - 0.55, body, 10, INK)


def add_plain_table(slide, x, y, w, h, rows, col_widths, header_fill=BLUE):
    row_h = h / len(rows)
    for r, row_values in enumerate(rows):
        cx = x
        for c, value in enumerate(row_values):
            cw = w * col_widths[c]
            cell = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(cx),
                Inches(y + r * row_h),
                Inches(cw),
                Inches(row_h),
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else RGBColor(248, 251, 255)
            cell.line.color.rgb = RGBColor(218, 228, 240)
            set_text(
                cell,
                str(value),
                10 if r == 0 else 9,
                WHITE if r == 0 else INK,
                r == 0,
                PP_ALIGN.CENTER if r == 0 else None,
            )
            cx += cw


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, width=2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_process_step(slide, x, y, w, num, title, body, color=BLUE):
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.42), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    add_text(slide, x, y + 0.08, 0.42, 0.16, str(num), 9, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, x + 0.55, y - 0.02, w - 0.55, 0.28, title, 13, color, True)
    add_text(slide, x + 0.55, y + 0.34, w - 0.55, 0.55, body, 9, INK)


def build_pptx() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    assets = prepare_assets()
    unicast = mean_by_protocol(ROOT / "results" / "exp_50n_unicast_pairs.csv")
    mixed = mean_by_protocol(ROOT / "results" / "exp_50n_mixed.csv")

    prs = Presentation(str(TEMPLATE))
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    clear_slides(prs)
    blank = prs.slide_layouts[6]

    slide_no = 1

    # 1 cover
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(assets["cover"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(4, 20, 43)
    overlay.fill.transparency = 28
    overlay.line.fill.background()
    s.shapes.add_picture(str(assets["logo"]), Inches(0.7), Inches(0.55), width=Inches(3.45))
    add_text(s, 0.82, 2.0, 8.1, 1.0, "LoRa Mesh\n路由仿真与协议对比", 36, WHITE, True)
    add_text(s, 0.86, 4.15, 8.0, 0.45, "从 managed flooding 到 route-cache source routing 的空口效率权衡", 15, RGBColor(220, 237, 255))
    add_text(s, 0.86, 5.0, 3.2, 0.32, "汇报人：BH4ME", 12, WHITE)
    add_text(s, 0.86, 5.38, 5.0, 0.32, "深圳大学电子与信息工程学院", 12, WHITE)
    slide_no += 1

    # 2 agenda
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "汇报路线", "先讲为什么需要 mesh，再看仿真器如何公平比较协议。", "00")
    agenda = ["问题背景：LoRaWAN 星型拓扑的边界", "仿真器：统一 PHY / 信道 / 碰撞模型", "基线协议：泛洪抑制 vs 路由缓存", "实验结果：可靠性、时延、空口占用的取舍", "下一步：面向 ToA / SNR 的路由改进"]
    for i, item in enumerate(agenda, 1):
        add_text(s, 1.05, 1.75 + i * 0.72, 0.7, 0.4, f"{i:02d}", 18, GREEN if i % 2 else BLUE, True, PP_ALIGN.CENTER)
        add_text(s, 1.85, 1.75 + i * 0.72, 8.6, 0.4, item, 18, DEEP, True)
    add_footer(s, assets, slide_no); slide_no += 1

    # 3 why mesh
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "为什么 LoRa 需要 Mesh", "LoRa 长距离、低功耗，但星型网关结构不总是适合野外和 P2P 场景。", "01")
    add_path_diagram(s)
    bullet_list(s, 0.95, 1.65, 5.2, [
        "扩大覆盖：节点之间接力，绕过遮挡与盲区",
        "降低基础设施依赖：不强依赖固定网关或互联网",
        "支持端到端通信：设备之间可直接形成多跳路径",
    ], 13)
    bullet_list(s, 7.65, 1.65, 4.2, [
        "代价：转发会消耗电量、CPU、内存和空口时间",
        "瓶颈：LoRa 低速率、半双工，碰撞与 duty cycle 敏感",
        "问题：该泛洪，还是先找路再单播？",
    ], 13)
    add_footer(s, assets, slide_no); slide_no += 1

    # 4 simulator
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "项目目标：一个公平的 LoRa Mesh 研究台架", "不是固件复刻，而是在同一拓扑、流量和物理层假设下比较路由行为。", "02")
    add_callout(s, 0.8, 1.55, 3.6, 1.25, "输入", "节点数、区域大小、流量类型、pair-count、种子、SF/BW/CR、发射功率、路径损耗", BLUE)
    add_callout(s, 4.85, 1.55, 3.65, 1.25, "内核", "事件驱动包级仿真：LoRa ToA、半双工、捕获效应、SNR 概率接收、碰撞失败", GREEN)
    add_callout(s, 8.95, 1.55, 3.25, 1.25, "输出", "PDR、覆盖率、时延、发包数、控制包、总空口时间、碰撞、重复接收、缓存命中", RED)
    add_text(s, 1.0, 3.65, 10.8, 0.52, "核心价值：把“路由策略”从随机拓扑与信道波动中剥离出来，用多 seed 平均值观察趋势。", 20, DEEP, True, PP_ALIGN.CENTER)
    add_metric(s, 1.4, 4.75, "50", "节点规模")
    add_metric(s, 3.7, 4.75, "1200s", "仿真时长")
    add_metric(s, 6.0, 4.75, "20", "随机种子")
    add_metric(s, 8.3, 4.75, "SF9", "固定 PHY")
    add_metric(s, 10.6, 4.75, "8", "固定会话对")
    add_footer(s, assets, slide_no); slide_no += 1

    # 5 protocols
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "两种基线：一个靠泛洪抑制，一个靠路径缓存", "同一仿真环境下比较 Meshtastic-like 与 MeshCore-like 的行为等价模型。", "03")
    add_callout(s, 0.9, 1.65, 5.2, 3.15, "Meshtastic-like：managed flooding", "每个新 DATA 包被延迟转发一次；节点听到重复副本后取消待转发任务。优点是鲁棒、覆盖好；缺点是重复接收和碰撞多，空口开销大。", BLUE)
    add_callout(s, 7.0, 1.65, 5.0, 3.15, "MeshCore-like：RREQ/RREP + source route", "首次单播泛洪路由请求，目的节点沿反向路径应答；后续同一对节点使用缓存路径。优点是重复流量时省空口；风险是发现阶段和路径质量会影响 PDR。", GREEN)
    add_text(s, 1.2, 5.35, 10.7, 0.55, "演讲抓手：LoRa mesh 的关键不是“谁一定更好”，而是“流量模式决定路由代价”。", 21, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 6 model
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "物理层与信道假设", "所有协议共享同一 SX1262-style profile，使比较集中在路由策略。", "04")
    s.shapes.add_picture(str(assets["chip"]), Inches(0.95), Inches(1.55), width=Inches(1.15))
    add_text(s, 2.35, 1.6, 9.3, 0.44, "LoRa time-on-air 是空口成本的共同货币", 22, DEEP, True)
    bullet_list(s, 1.25, 2.45, 10.0, [
        "Log-distance path loss + log-normal shadowing：生成虚拟 RSSI/SNR",
        "固定 SF/BW/CR：当前结果控制 PHY 变量，便于看路由差异",
        "Half-duplex radio：正在发射的节点不能接收",
        "Collision + capture threshold：强信号可捕获，弱信号失败",
        "SNR margin -> probabilistic reception：链路不是硬阈值，而是概率成功",
    ], 15, INK, 0.52)
    add_text(s, 1.05, 5.75, 10.9, 0.5, "来自文献的启发：ToA / 多 SF-aware 路由在异构拓扑中有机会提升 PDR、公平性与时延表现。", 16, BLUE, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 7 unicast result
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "重复单播：缓存路径显著节省空口，但可靠性下降", "场景：50 nodes / 1200s / unicast / pair-count 8 / 20 seeds。", "05")
    add_bar_chart(
        s, 0.85, 1.55, 4.0, 2.4, "可靠性与时延",
        ["PDR", "Avg delay"],
        "Meshtastic-like",
        [unicast["meshtastic-like"]["unicast_pdr"], unicast["meshtastic-like"]["avg_delay_s"]],
        "MeshCore-like",
        [unicast["meshcore-like"]["unicast_pdr"], unicast["meshcore-like"]["avg_delay_s"]],
    )
    add_bar_chart(
        s, 5.1, 1.55, 3.5, 2.4, "总发送次数",
        ["TX count"],
        "Meshtastic-like",
        [unicast["meshtastic-like"]["tx_count"]],
        "MeshCore-like",
        [unicast["meshcore-like"]["tx_count"]],
    )
    add_bar_chart(
        s, 8.85, 1.55, 3.35, 2.4, "总空口时间(s)",
        ["airtime"],
        "Meshtastic-like",
        [unicast["meshtastic-like"]["total_airtime_s"]],
        "MeshCore-like",
        [unicast["meshcore-like"]["total_airtime_s"]],
    )
    add_metric(s, 1.25, 4.65, "70%", "MeshCore-like 发包数下降", GREEN)
    add_metric(s, 3.7, 4.65, "70%", "总空口时间下降", GREEN)
    add_metric(s, 6.15, 4.65, "77%", "碰撞失败下降", GREEN)
    add_metric(s, 8.6, 4.65, "0.83", "MeshCore-like PDR", RED)
    add_text(s, 1.0, 6.1, 11.2, 0.35, "解读：重复会话让缓存收益出现，但路径发现与固定源路由让部分链路上的失败更集中。", 15, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 8 mixed result
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "混合流量：广播覆盖保持高位，单播可靠性成为代价", "场景：50 nodes / 1200s / mixed / pair-count 8 / 20 seeds。", "06")
    add_bar_chart(
        s, 0.9, 1.5, 4.3, 2.8, "PDR / 覆盖率",
        ["Unicast PDR", "Broadcast coverage"],
        "Meshtastic-like",
        [mixed["meshtastic-like"]["unicast_pdr"], mixed["meshtastic-like"]["broadcast_coverage"]],
        "MeshCore-like",
        [mixed["meshcore-like"]["unicast_pdr"], mixed["meshcore-like"]["broadcast_coverage"]],
    )
    add_bar_chart(
        s, 5.55, 1.5, 3.0, 2.8, "总空口时间(s)",
        ["airtime"],
        "Meshtastic-like",
        [mixed["meshtastic-like"]["total_airtime_s"]],
        "MeshCore-like",
        [mixed["meshcore-like"]["total_airtime_s"]],
    )
    add_bar_chart(
        s, 8.95, 1.5, 3.05, 2.8, "碰撞失败",
        ["collision"],
        "Meshtastic-like",
        [mixed["meshtastic-like"]["collision_fail"]],
        "MeshCore-like",
        [mixed["meshcore-like"]["collision_fail"]],
    )
    add_text(s, 1.1, 4.85, 10.9, 0.55, "混合场景中，MeshCore-like 的广播覆盖略高，但 unicast PDR 从 0.95 降到 0.64；这说明广播 fallback 与路径单播的目标函数不同。", 16, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 9 interpretation
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "结论不是选边站，而是识别流量结构", "不同流量下，可靠性、时延和空口效率形成三角取舍。", "07")
    add_text(s, 1.15, 1.6, 3.2, 0.35, "可靠性", 18, BLUE, True, PP_ALIGN.CENTER)
    add_text(s, 5.1, 1.6, 3.2, 0.35, "空口效率", 18, GREEN, True, PP_ALIGN.CENTER)
    add_text(s, 8.95, 1.6, 3.2, 0.35, "适用流量", 18, RED, True, PP_ALIGN.CENTER)
    bullet_list(s, 0.95, 2.25, 3.5, ["Managed flooding 更稳", "重复副本增加成功概率", "代价是碰撞与冗余"], 13, INK, 0.55)
    bullet_list(s, 4.9, 2.25, 3.5, ["Route cache 更省", "路径建立后只走必要节点", "代价是路径质量敏感"], 13, INK, 0.55)
    bullet_list(s, 8.75, 2.25, 3.5, ["重复单播：缓存有价值", "广播/组播：泛洪仍重要", "高负载：需要拥塞感知"], 13, INK, 0.55)
    add_text(s, 1.2, 5.35, 10.7, 0.62, "因此下一版路由不该只是“缓存路径”，而应把 ToA、SNR margin、拥塞与路径新鲜度放进同一个成本函数。", 20, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 10 next
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "下一步：从基线比较走向自适应路由", "把仿真器变成协议设计迭代平台。", "08")
    add_callout(s, 0.85, 1.5, 3.55, 3.0, "1. 加入链路代价", "把 ToA、SNR margin、重传风险、节点拥塞转换为路径成本，避免只按 hop count 选路。", BLUE)
    add_callout(s, 4.85, 1.5, 3.55, 3.0, "2. 扩展 PHY 维度", "支持多 SF / 多信道，把 LoRa 的准正交特性纳入路由与调度。", GREEN)
    add_callout(s, 8.85, 1.5, 3.35, 3.0, "3. 做实证闭环", "用真实 SX1262 RSSI/SNR 日志校准传播模型，并对比实测与仿真趋势。", RED)
    add_text(s, 1.0, 5.35, 11.1, 0.62, "目标：在不牺牲太多 PDR 的前提下，把空口占用和碰撞压下来。", 22, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 11 thanks
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(assets["campus"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(255, 255, 255)
    overlay.fill.transparency = 10
    overlay.line.fill.background()
    s.shapes.add_picture(str(assets["mark"]), Inches(0.82), Inches(0.65), width=Inches(1.08))
    add_text(s, 2.1, 2.25, 8.9, 0.9, "谢谢聆听", 39, DEEP, True, PP_ALIGN.CENTER)
    add_text(s, 2.1, 3.38, 8.9, 0.5, "欢迎交流 LoRa Mesh 路由、仿真参数与下一步实验设计", 17, BLUE, True, PP_ALIGN.CENTER)
    add_text(s, 2.1, 5.45, 8.9, 0.28, "LoRa Mesh Routing Simulator · MIT License", 10, MUTED, False, PP_ALIGN.CENTER)

    prs.save(PPTX_OUT)
    render_previews(assets, unicast, mixed)


def build_pptx() -> None:
    """Build the refined version used for delivery."""

    OUT_DIR.mkdir(exist_ok=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    assets = prepare_assets()
    unicast = mean_by_protocol(ROOT / "results" / "exp_50n_unicast_pairs.csv")
    mixed = mean_by_protocol(ROOT / "results" / "exp_50n_mixed.csv")
    unicast_detail = summary_by_protocol(ROOT / "results" / "exp_50n_unicast_pairs.csv")
    mixed_detail = summary_by_protocol(ROOT / "results" / "exp_50n_mixed.csv")

    prs = Presentation(str(TEMPLATE))
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    clear_slides(prs)
    blank = prs.slide_layouts[6]
    slide_no = 1

    # 1. Cover
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(assets["cover"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(3, 19, 42)
    overlay.fill.transparency = 22
    overlay.line.fill.background()
    s.shapes.add_picture(str(assets["logo"]), Inches(0.7), Inches(0.54), width=Inches(3.42))
    add_text(s, 0.84, 1.88, 9.2, 1.25, "LoRa Mesh 路由仿真\n与协议取舍研究", 34, WHITE, True)
    add_text(s, 0.88, 4.1, 8.6, 0.55, "面向低速率、半双工无线网络的 PDR / 时延 / 空口效率比较", 15, RGBColor(224, 238, 255))
    add_text(s, 0.9, 5.15, 5.2, 0.34, "汇报人：BH4ME｜深圳大学电子与信息工程学院", 11, WHITE)
    slide_no += 1

    # 2. Project intro
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "先用一句话理解", "正在研究的是：LoRa 节点帮彼此转发消息时，怎样做到既送得到、又少占无线资源。", "00")
    add_text(s, 1.0, 1.48, 10.8, 0.82, "LoRa Mesh = 不完全依赖网关，让节点之间一跳一跳接力，把消息送到更远的地方。", 21, DEEP, True, PP_ALIGN.CENTER)
    add_callout(s, 0.95, 2.65, 3.25, 2.15, "关系一", "转发方法会影响结果：是大家都帮忙转发，还是先找一条路再沿路发送？", BLUE)
    add_callout(s, 4.95, 2.65, 3.25, 2.15, "关系二", "更稳往往更费资源：多发几份更容易送到，但也更容易占用空口、造成碰撞。", GREEN)
    add_callout(s, 8.95, 2.65, 3.05, 2.15, "怎么研究", "让两种方法在同样的节点、信道和流量下跑仿真，再比较送达率、延迟和发包数量。", RED)
    add_text(s, 1.15, 5.55, 10.6, 0.42, "所以这次汇报先不追求复杂公式，重点看清楚：省资源和送得稳之间怎么平衡。", 16, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 3. Executive message
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "核心结论先说人话", "缓存路线能少发很多包，但如果路线本身不够好，有些消息就更容易送不到。", "01")
    add_text(s, 1.0, 1.65, 10.9, 0.95, "LoRa Mesh 的难点不是“用哪一个协议名字”，\n而是在低速率共享空口里，找到“少发包”和“送得稳”的平衡点。", 24, DEEP, True, PP_ALIGN.CENTER)
    add_metric(s, 1.45, 3.35, f"{pct_drop(unicast['meshtastic-like']['tx_count'], unicast['meshcore-like']['tx_count']):.0%}", "重复单播 TX 下降", GREEN)
    add_metric(s, 3.85, 3.35, f"{pct_drop(unicast['meshtastic-like']['total_airtime_s'], unicast['meshcore-like']['total_airtime_s']):.0%}", "重复单播空口下降", GREEN)
    add_metric(s, 6.25, 3.35, f"{unicast['meshtastic-like']['unicast_pdr']:.2f}", "泛洪 PDR", BLUE)
    add_metric(s, 8.65, 3.35, f"{unicast['meshcore-like']['unicast_pdr']:.2f}", "缓存路径 PDR", RED)
    add_text(s, 1.25, 5.35, 10.4, 0.5, "后续优化方向：不要只记住“有一条路”，还要判断这条路现在好不好、忙不忙、稳不稳。", 17, BLUE, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 4. Agenda
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "目录", "下面每一部分都对应后面的页面标题，避免听的时候找不到位置。", "TOC")
    agenda = [
        ("00", "先用一句话理解", "LoRa Mesh 在研究什么问题"),
        ("01", "核心结论", "少发包和送得稳之间要平衡"),
        ("02", "背景与问题", "为什么需要 LoRa Mesh、主要看哪两件事"),
        ("03", "仿真内容", "怎么公平比较、仿真怎么跑、参数怎么设"),
        ("04", "两种转发思路", "大家帮忙转 vs 先找路再沿路转"),
        ("05", "仿真结果", "重复单播和混合流量的对比"),
        ("06", "原因与改进", "为什么会丢、下一步怎么改"),
    ]
    for i, (num, head, body) in enumerate(agenda):
        y = 1.45 + i * 0.68
        add_text(s, 1.02, y, 0.72, 0.35, num, 17, GREEN if i % 2 == 0 else BLUE, True, PP_ALIGN.CENTER)
        add_text(s, 1.95, y - 0.02, 2.0, 0.32, head, 17, DEEP, True)
        add_text(s, 3.82, y + 0.02, 7.6, 0.28, body, 12, MUTED)
    add_footer(s, assets, slide_no); slide_no += 1

    # 5. Problem framing
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "为什么需要 LoRa Mesh", "LoRa 传得远，但只靠网关时，覆盖范围和部署条件还是会受限制。", "02")
    add_callout(s, 0.85, 1.55, 3.45, 1.55, "星型拓扑的限制", "网关覆盖决定网络边界；复杂地形、野外部署、断网场景会让端节点之间难以直接通信。", BLUE)
    add_callout(s, 4.95, 1.55, 3.45, 1.55, "Mesh 的收益", "节点可接力转发，扩大覆盖并支持设备间多跳通信，适合巡检、应急、低基础设施场景。", GREEN)
    add_callout(s, 9.05, 1.55, 3.1, 1.55, "新的代价", "每一次转发都消耗电量与 ToA；低速率半双工让碰撞、冗余和延迟变得敏感。", RED)
    add_path_diagram(s)
    add_text(s, 1.05, 5.6, 11.1, 0.48, "因此，路由协议的任务不是盲目减少跳数，而是在共享空口上选择“值得发送”的下一跳。", 18, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 6. Research question
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "主要看两件事", "不是先追求复杂协议，而是先看清楚两种基本转发思路的差别。", "02")
    add_text(s, 1.1, 1.6, 4.8, 0.52, "问题 1", 16, BLUE, True)
    add_text(s, 1.1, 2.0, 4.8, 0.92, "大家都帮忙转发，消息是不是更容易送到？", 22, DEEP, True)
    add_text(s, 7.05, 1.6, 4.8, 0.52, "问题 2", 16, GREEN, True)
    add_text(s, 7.05, 2.0, 4.8, 0.92, "先找好路线再发送，能省多少资源？", 22, DEEP, True)
    add_plain_table(
        s, 1.1, 3.45, 10.9, 1.85,
        [
            ["比较维度", "关注指标", "原因"],
            ["可靠性", "unicast PDR / broadcast coverage", "LoRa mesh 首先要把小包送到"],
            ["实时性", "avg delay", "多跳与路由发现会增加端到端等待"],
            ["空口成本", "tx_count / total_airtime / collision_fail", "LoRa 低速率，共享空口是核心稀缺资源"],
        ],
        [0.22, 0.35, 0.43],
    )
    add_footer(s, assets, slide_no); slide_no += 1

    # 7. Simulator architecture
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "怎么比较才公平", "让两种方法在同样的节点、无线环境和流量里运行。", "03")
    stages = [
        ("Topology", "随机节点坐标\nrouter / repeater role"),
        ("Radio", "SF9 / 125 kHz / CR 4/5\nLoRa time-on-air"),
        ("Channel", "路径损耗 + 阴影衰落\n半双工 + 捕获效应"),
        ("Traffic", "unicast / broadcast / mixed\n固定会话对 pair-count"),
        ("Metrics", "PDR / delay / airtime\ncollision / duplicate / cache"),
    ]
    for i, (head, body) in enumerate(stages):
        x = 0.75 + i * 2.45
        color = [BLUE, GREEN, BLUE, GREEN, RED][i]
        add_callout(s, x, 1.75, 2.0, 2.1, head, body, color)
        if i < len(stages) - 1:
            add_arrow(s, x + 2.06, 2.8, x + 2.36, 2.8, RGBColor(126, 156, 191), 2)
    add_text(s, 1.0, 4.75, 11.1, 0.72, "公平性来自三个共享：同一拓扑生成、同一传播/碰撞模型、同一 traffic schedule。这样协议差异不会被随机信道差异掩盖。", 18, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 8. Simulation flow
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "仿真到底怎么跑", "把真实无线网络简化成一套可重复的事件流程：谁发、谁收、是否碰撞、最后统计结果。", "03")
    add_callout(s, 0.85, 1.45, 3.2, 1.75, "输入", "节点数量和位置、无线参数、流量类型、随机种子、最大跳数、仿真时长。", BLUE)
    add_callout(s, 5.05, 1.45, 3.2, 1.75, "过程", "按时间顺序处理事件：产生消息、开始发射、计算接收、判断碰撞、调用路由协议转发。", GREEN)
    add_callout(s, 9.25, 1.45, 2.95, 1.75, "输出", "每个协议输出 CSV：送达率、延迟、发包数、空口时间、碰撞失败等指标。", RED)
    add_process_step(s, 0.95, 3.65, 10.8, 1, "生成网络", "在 3000 m 区域内随机放置 50 个节点，固定一组无线参数。", BLUE)
    add_process_step(s, 0.95, 4.45, 10.8, 2, "生成流量", "按 6 flows/min 产生消息；单播场景固定 8 对节点反复通信。", GREEN)
    add_process_step(s, 0.95, 5.25, 10.8, 3, "重复多次", "同一场景跑 20 个随机种子，最后比较平均结果，而不是看单次偶然结果。", RED)
    add_footer(s, assets, slide_no); slide_no += 1

    # 9. Protocol mechanism
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "两种转发思路", "一种是“多几个人帮忙转”，一种是“先找路，再沿路转”。", "04")
    add_text(s, 0.95, 1.55, 5.25, 0.38, "Meshtastic-like：managed flooding", 18, BLUE, True)
    add_process_step(s, 1.05, 2.05, 4.5, 1, "DATA 广播", "源节点发送，邻居收到后进入延迟转发。", BLUE)
    add_process_step(s, 1.05, 3.05, 4.5, 2, "延迟 + 抑制", "若提前听到其他副本，取消自己的转发。", BLUE)
    add_process_step(s, 1.05, 4.05, 4.5, 3, "以冗余换鲁棒", "重复副本提升到达概率，但增加碰撞与 ToA。", BLUE)
    add_text(s, 7.0, 1.55, 5.0, 0.38, "MeshCore-like：RREQ / RREP / source route", 18, GREEN, True)
    add_process_step(s, 7.05, 2.05, 4.6, 1, "首次找路", "RREQ 泛洪，目的节点记录反向路径。", GREEN)
    add_process_step(s, 7.05, 3.05, 4.6, 2, "路径应答", "RREP 沿反向路径返回，源节点写入缓存。", GREEN)
    add_process_step(s, 7.05, 4.05, 4.6, 3, "按路径单播", "后续流量只让路径上的下一跳转发。", GREEN)
    add_text(s, 1.2, 5.75, 10.7, 0.4, "这页演讲时可以用一句话带过：一种方法更稳但更吵，另一种更省但更怕路线不好。", 15, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 10. Experiment setup
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "仿真场景和参数", "先固定无线参数，把注意力放在转发方法本身。", "03")
    add_plain_table(
        s, 0.9, 1.5, 5.4, 4.2,
        [
            ["参数", "取值"],
            ["节点数 / 区域", "50 nodes / 3000 m"],
            ["仿真时长", "1200 s"],
            ["流量速率", "6 flows/min"],
            ["重复会话", "pair-count = 8"],
            ["随机种子", "20 seeds"],
            ["最大跳数", "7 hops"],
        ],
        [0.42, 0.58],
    )
    add_plain_table(
        s, 6.9, 1.5, 4.9, 4.2,
        [
            ["PHY / 信道", "取值"],
            ["SF / BW / CR", "SF9 / 125 kHz / 4/5"],
            ["payload", "32 bytes"],
            ["TX power", "17 dBm"],
            ["path loss exp.", "2.7"],
            ["shadow sigma", "4 dB"],
            ["capture threshold", "6 dB"],
        ],
        [0.45, 0.55],
        GREEN,
    )
    add_footer(s, assets, slide_no); slide_no += 1

    # 11. Unicast results
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "结果一：重复单播下，路径缓存把空口压力大幅压低", "但 PDR 从 0.95 降到 0.83，说明缓存路径还缺少链路质量判断。", "05")
    add_bar_chart(
        s, 0.75, 1.45, 3.55, 2.45, "PDR 与平均时延",
        ["PDR", "delay(s)"],
        "Meshtastic-like",
        [unicast["meshtastic-like"]["unicast_pdr"], unicast["meshtastic-like"]["avg_delay_s"]],
        "MeshCore-like",
        [unicast["meshcore-like"]["unicast_pdr"], unicast["meshcore-like"]["avg_delay_s"]],
    )
    add_bar_chart(
        s, 4.72, 1.45, 3.2, 2.45, "发送次数",
        ["total TX", "data TX"],
        "Meshtastic-like",
        [unicast["meshtastic-like"]["tx_count"], unicast["meshtastic-like"]["data_tx"]],
        "MeshCore-like",
        [unicast["meshcore-like"]["tx_count"], unicast["meshcore-like"]["data_tx"]],
    )
    add_bar_chart(
        s, 8.35, 1.45, 3.55, 2.45, "空口与碰撞",
        ["airtime(s)", "collision/100"],
        "Meshtastic-like",
        [unicast["meshtastic-like"]["total_airtime_s"], unicast["meshtastic-like"]["collision_fail"] / 100],
        "MeshCore-like",
        [unicast["meshcore-like"]["total_airtime_s"], unicast["meshcore-like"]["collision_fail"] / 100],
    )
    add_metric(s, 1.2, 4.65, f"{pct_drop(unicast['meshtastic-like']['tx_count'], unicast['meshcore-like']['tx_count']):.0%}", "TX 下降", GREEN)
    add_metric(s, 3.55, 4.65, f"{pct_drop(unicast['meshtastic-like']['total_airtime_s'], unicast['meshcore-like']['total_airtime_s']):.0%}", "空口下降", GREEN)
    add_metric(s, 5.9, 4.65, f"{pct_drop(unicast['meshtastic-like']['collision_fail'], unicast['meshcore-like']['collision_fail']):.0%}", "碰撞失败下降", GREEN)
    add_metric(s, 8.25, 4.65, f"{unicast_detail['meshcore-like']['route_cache_hits']['mean']:.0f}", "平均缓存命中", BLUE)
    add_text(s, 1.1, 6.05, 10.9, 0.32, "讲法：MeshCore-like 的控制包很多，但换来后续 data TX 极少；问题是路径一旦不稳，收益会转化为丢包。", 13, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 12. Mixed results
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "结果二：混合流量中，广播覆盖仍高，但单播可靠性掉得更多", "广播 fallback 与路径单播的目标不一致，导致协议表现分化。", "05")
    add_bar_chart(
        s, 0.8, 1.45, 3.75, 2.7, "可靠性",
        ["unicast PDR", "broadcast cov."],
        "Meshtastic-like",
        [mixed["meshtastic-like"]["unicast_pdr"], mixed["meshtastic-like"]["broadcast_coverage"]],
        "MeshCore-like",
        [mixed["meshcore-like"]["unicast_pdr"], mixed["meshcore-like"]["broadcast_coverage"]],
    )
    add_bar_chart(
        s, 4.95, 1.45, 3.35, 2.7, "总空口时间(s)",
        ["airtime"],
        "Meshtastic-like",
        [mixed["meshtastic-like"]["total_airtime_s"]],
        "MeshCore-like",
        [mixed["meshcore-like"]["total_airtime_s"]],
    )
    add_bar_chart(
        s, 8.7, 1.45, 3.25, 2.7, "碰撞失败",
        ["collision"],
        "Meshtastic-like",
        [mixed["meshtastic-like"]["collision_fail"]],
        "MeshCore-like",
        [mixed["meshcore-like"]["collision_fail"]],
    )
    add_text(s, 1.0, 4.75, 3.3, 0.55, f"{mixed['meshcore-like']['broadcast_coverage']:.2f}", 29, GREEN, True, PP_ALIGN.CENTER)
    add_text(s, 1.0, 5.35, 3.3, 0.3, "MeshCore-like 广播覆盖", 10, MUTED, False, PP_ALIGN.CENTER)
    add_text(s, 4.95, 4.75, 3.3, 0.55, f"{mixed['meshcore-like']['unicast_pdr']:.2f}", 29, RED, True, PP_ALIGN.CENTER)
    add_text(s, 4.95, 5.35, 3.3, 0.3, "MeshCore-like 单播 PDR", 10, MUTED, False, PP_ALIGN.CENTER)
    add_text(s, 8.55, 4.85, 3.4, 0.55, f"{pct_drop(mixed['meshtastic-like']['collision_fail'], mixed['meshcore-like']['collision_fail']):.0%}", 29, GREEN, True, PP_ALIGN.CENTER)
    add_text(s, 8.55, 5.42, 3.4, 0.3, "碰撞失败下降", 10, MUTED, False, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 13. Why PDR drops
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "为什么少发包反而更容易丢", "缓存路线省掉了重复发送，也减少了“碰巧从别的路径送到”的机会。", "06")
    add_callout(s, 0.9, 1.55, 3.35, 2.55, "1. 路径发现有成本", "RREQ/RREP 会占用空口；发现失败或应答失败时，首包会被延迟甚至无法送达。", BLUE)
    add_callout(s, 4.85, 1.55, 3.35, 2.55, "2. 源路由路径脆弱", "后续数据只走缓存路径；路径上任意关键链路质量差，PDR 就被单点放大影响。", GREEN)
    add_callout(s, 8.8, 1.55, 3.25, 2.55, "3. 固定 PHY 不自适应", "当前所有节点固定 SF9，尚未利用多 SF 或链路 ToA 差异来避开拥塞路径。", RED)
    add_text(s, 1.05, 4.85, 11.1, 0.7, "改进方向很明确：路径缓存仍然有价值，但缓存项必须带有质量、拥塞和过期策略，而不只是“有没有路径”。", 20, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 14. Proposed improvement
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "下一步：让节点会判断路线好不好", "目标是在保持送达率的同时，减少无效转发和碰撞。", "06")
    add_text(s, 0.95, 1.65, 11.3, 0.52, "下一版路线选择不只看“有没有路”，还要看四个因素", 18, BLUE, True)
    add_text(s, 1.05, 2.25, 11.0, 0.62, "路线是否短、链路是否稳、节点是否忙、这条路线是否已经过期", 23, DEEP, True, PP_ALIGN.CENTER)
    add_plain_table(
        s, 1.0, 3.35, 11.0, 2.05,
        [
            ["判断因素", "通俗含义", "希望达到的效果"],
            ["发一次要多久", "这条链路占用无线空口的时间", "避开又慢又占资源的路径"],
            ["链路稳不稳", "RSSI/SNR 反映这条链路是否容易丢包", "不要只选省包但不可靠的路线"],
            ["节点忙不忙", "中继节点是否排队、碰撞多、频道忙", "绕开热点节点"],
            ["路线新不新", "缓存路线有没有过期，是否需要重新探测", "避免一直使用旧路线"],
        ],
        [0.22, 0.4, 0.38],
    )
    add_footer(s, assets, slide_no); slide_no += 1

    # 15. Roadmap
    s = prs.slides.add_slide(blank)
    add_soft_bg(s, assets)
    add_title(s, "后续工作：让仿真与真实硬件闭环", "把当前 baseline 变成协议迭代流程。", "06")
    add_process_step(s, 0.9, 1.65, 10.8, 1, "仿真侧扩展", "加入多 SF / 多信道、链路质量缓存、路由老化、拥塞感知转发策略。", BLUE)
    add_process_step(s, 0.9, 2.85, 10.8, 2, "实验侧校准", "采集 SX1262 真实 RSSI/SNR/丢包日志，校准 path loss、shadowing 与 PRR 曲线。", GREEN)
    add_process_step(s, 0.9, 4.05, 10.8, 3, "评估侧补强", "增加置信区间、参数扫频、节点密度变化、移动/失效节点和不同 traffic mix。", RED)
    add_text(s, 1.15, 5.75, 10.75, 0.42, "最终交付：一套可复现实验脚本 + 一组清楚说明协议边界的图表。", 17, DEEP, True, PP_ALIGN.CENTER)
    add_footer(s, assets, slide_no); slide_no += 1

    # 16. Thanks
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(assets["campus"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(255, 255, 255)
    overlay.fill.transparency = 8
    overlay.line.fill.background()
    s.shapes.add_picture(str(assets["mark"]), Inches(0.82), Inches(0.65), width=Inches(1.08))
    add_text(s, 2.1, 2.22, 8.9, 0.9, "谢谢聆听", 39, DEEP, True, PP_ALIGN.CENTER)
    add_text(s, 2.1, 3.32, 8.9, 0.5, "欢迎交流 LoRa Mesh 路由、仿真参数与下一步实验设计", 17, BLUE, True, PP_ALIGN.CENTER)
    add_text(s, 2.1, 5.45, 8.9, 0.28, "LoRa Mesh Routing Simulator · baseline results from 20 seeds", 10, MUTED, False, PP_ALIGN.CENTER)

    prs.save(PPTX_OUT)
    render_previews_v2(assets, unicast, mixed)


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(PIL_FONT, size=size, index=0)


def draw_center(draw: ImageDraw.ImageDraw, xy, text, font, fill):
    x, y, w, h = xy
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=8)
    draw.multiline_text((x + (w - (bbox[2] - bbox[0])) / 2, y + (h - (bbox[3] - bbox[1])) / 2), text, font=font, fill=fill, align="center", spacing=8)


def wrap_text(text: str, limit: int) -> str:
    lines: list[str] = []
    current = ""
    for ch in text:
        current += ch
        if len(current) >= limit and ch in "，。：； /":
            lines.append(current.strip())
            current = ""
    if current:
        lines.append(current.strip())
    return "\n".join(lines)


def preview_base(assets):
    im = Image.new("RGB", (1920, 1080), (249, 252, 255))
    try:
        bg = Image.open(assets["light_circuit"]).convert("RGBA")
        bg.thumbnail((780, 420))
        im.paste(bg, (1120, 40), bg)
    except Exception:
        pass
    return im


def draw_title(im, title, subtitle="", section=""):
    d = ImageDraw.Draw(im)
    d.text((112, 62), title, font=pil_font(46, True), fill=(17, 38, 74))
    if subtitle:
        d.text((115, 138), subtitle, font=pil_font(22), fill=(95, 111, 132))
    if section:
        d.text((1650, 70), section, font=pil_font(30, True), fill=(20, 104, 184))
    d.rounded_rectangle((112, 181, 230, 188), radius=4, fill=(111, 186, 45))


def draw_footer(im, n):
    d = ImageDraw.Draw(im)
    d.text((112, 1000), "深圳大学电子与信息工程学院", font=pil_font(15), fill=(95, 111, 132))
    d.text((1720, 998), f"{n:02d}", font=pil_font(18, True), fill=(20, 104, 184))


def render_previews(assets, unicast, mixed) -> None:
    for p in PREVIEW_DIR.glob("*.png"):
        p.unlink()

    slides: list[Image.Image] = []

    im = Image.open(assets["cover"]).convert("RGB").resize((1920, 1080))
    shade = Image.new("RGBA", im.size, (4, 20, 43, 95))
    im = Image.alpha_composite(im.convert("RGBA"), shade).convert("RGB")
    d = ImageDraw.Draw(im)
    d.multiline_text((118, 295), "LoRa Mesh\n路由仿真与协议对比", font=pil_font(72, True), fill=(255, 255, 255), spacing=18)
    d.text((124, 600), "从 managed flooding 到 route-cache source routing 的空口效率权衡", font=pil_font(29), fill=(220, 237, 255))
    d.text((124, 720), "汇报人：BH4ME", font=pil_font(24), fill=(255, 255, 255))
    slides.append(im)

    titles = [
        ("汇报路线", "先讲为什么需要 mesh，再看仿真器如何公平比较协议。"),
        ("为什么 LoRa 需要 Mesh", "LoRa 长距离、低功耗，但星型网关结构不总是适合野外和 P2P 场景。"),
        ("项目目标：一个公平的 LoRa Mesh 研究台架", "不是固件复刻，而是在同一拓扑、流量和物理层假设下比较路由行为。"),
        ("两种基线：一个靠泛洪抑制，一个靠路径缓存", "同一仿真环境下比较 Meshtastic-like 与 MeshCore-like 的行为等价模型。"),
        ("物理层与信道假设", "所有协议共享同一 SX1262-style profile，使比较集中在路由策略。"),
        ("重复单播：缓存路径显著节省空口，但可靠性下降", "50 nodes / 1200s / unicast / pair-count 8 / 20 seeds。"),
        ("混合流量：广播覆盖保持高位，单播可靠性成为代价", "50 nodes / 1200s / mixed / pair-count 8 / 20 seeds。"),
        ("结论不是选边站，而是识别流量结构", "不同流量下，可靠性、时延和空口效率形成三角取舍。"),
        ("下一步：从基线比较走向自适应路由", "把仿真器变成协议设计迭代平台。"),
    ]
    for idx, (title, subtitle) in enumerate(titles, 2):
        im = preview_base(assets)
        draw_title(im, title, subtitle, f"{idx-2:02d}")
        d = ImageDraw.Draw(im)
        if idx == 2:
            items = ["01 问题背景：LoRaWAN 星型拓扑的边界", "02 仿真器：统一 PHY / 信道 / 碰撞模型", "03 基线协议：泛洪抑制 vs 路由缓存", "04 实验结果：可靠性、时延、空口占用的取舍", "05 下一步：面向 ToA / SNR 的路由改进"]
            for j, item in enumerate(items):
                d.text((180, 280 + j * 105), item, font=pil_font(34, True), fill=(17, 38, 74))
        elif idx == 3:
            left = ["扩大覆盖：节点之间接力，绕过遮挡与盲区", "降低基础设施依赖：不强依赖固定网关或互联网", "支持端到端通信：设备之间可直接形成多跳路径"]
            right = ["代价：转发会消耗电量、CPU、内存和空口时间", "瓶颈：LoRa 低速率、半双工，碰撞敏感", "问题：该泛洪，还是先找路再单播？"]
            for j, item in enumerate(left):
                d.ellipse((170, 275 + j * 78, 188, 293 + j * 78), fill=(111, 186, 45))
                d.text((212, 262 + j * 78), item, font=pil_font(27), fill=(31, 41, 55))
            for j, item in enumerate(right):
                d.ellipse((1110, 275 + j * 78, 1128, 293 + j * 78), fill=(20, 104, 184))
                d.text((1152, 262 + j * 78), item, font=pil_font(27), fill=(31, 41, 55))
            nodes = [(440, 650), (620, 540), (820, 630), (1010, 520), (1200, 640), (1400, 560)]
            for a, b in zip(nodes, nodes[1:]):
                d.line((a[0], a[1], b[0], b[1]), fill=(126, 156, 191), width=5)
            for k, (x, y) in enumerate(nodes):
                fill = (20, 104, 184) if k in (0, len(nodes) - 1) else (111, 186, 45)
                d.ellipse((x - 22, y - 22, x + 22, y + 22), fill=fill, outline=(255, 255, 255), width=4)
            draw_center(d, (250, 735, 1420, 80), "多跳 mesh：用节点转发补足网关覆盖、地形遮挡与 P2P 通信需求", pil_font(30, True), (17, 38, 74))
        elif idx == 4:
            boxes = [
                ("输入", "节点数、区域大小、流量类型、pair-count、种子、SF/BW/CR、发射功率、路径损耗", (20, 104, 184)),
                ("内核", "事件驱动包级仿真：LoRa ToA、半双工、捕获效应、SNR 概率接收、碰撞失败", (111, 186, 45)),
                ("输出", "PDR、覆盖率、时延、发包数、控制包、总空口时间、碰撞、重复接收、缓存命中", (194, 48, 58)),
            ]
            for j, (head, body, color) in enumerate(boxes):
                x = 145 + j * 585
                d.rounded_rectangle((x, 245, x + 500, 430), radius=18, fill=(244, 249, 255), outline=(214, 228, 244), width=2)
                d.text((x + 28, 268), head, font=pil_font(31, True), fill=color)
                d.multiline_text((x + 28, 325), wrap_text(body, 22), font=pil_font(20), fill=(31, 41, 55), spacing=6)
            draw_center(d, (190, 530, 1540, 85), "核心价值：把“路由策略”从随机拓扑与信道波动中剥离出来，用多 seed 平均值观察趋势。", pil_font(36, True), (17, 38, 74))
            for j, (num, lab) in enumerate([("50", "节点规模"), ("1200s", "仿真时长"), ("20", "随机种子"), ("SF9", "固定 PHY"), ("8", "固定会话对")]):
                draw_center(d, (250 + j * 300, 700, 220, 65), num, pil_font(48, True), (20, 104, 184))
                draw_center(d, (250 + j * 300, 775, 220, 45), lab, pil_font(20), (95, 111, 132))
        elif idx == 5:
            d.rounded_rectangle((150, 250, 870, 680), radius=22, fill=(244, 249, 255), outline=(214, 228, 244), width=2)
            d.rounded_rectangle((1030, 250, 1740, 680), radius=22, fill=(244, 249, 255), outline=(214, 228, 244), width=2)
            d.text((205, 288), "Meshtastic-like：managed flooding", font=pil_font(32, True), fill=(20, 104, 184))
            d.multiline_text((205, 360), wrap_text("每个新 DATA 包被延迟转发一次；节点听到重复副本后取消待转发任务。优点是鲁棒、覆盖好；缺点是重复接收和碰撞多，空口开销大。", 25), font=pil_font(25), fill=(31, 41, 55), spacing=10)
            d.text((1085, 288), "MeshCore-like：RREQ/RREP + source route", font=pil_font(32, True), fill=(111, 186, 45))
            d.multiline_text((1085, 360), wrap_text("首次单播泛洪路由请求，目的节点沿反向路径应答；后续同一对节点使用缓存路径。优点是重复流量时省空口；风险是路径质量会影响 PDR。", 25), font=pil_font(25), fill=(31, 41, 55), spacing=10)
            draw_center(d, (210, 760, 1500, 70), "演讲抓手：LoRa mesh 的关键不是“谁一定更好”，而是“流量模式决定路由代价”。", pil_font(34, True), (17, 38, 74))
        elif idx == 6:
            d.text((345, 242), "LoRa time-on-air 是空口成本的共同货币", font=pil_font(41, True), fill=(17, 38, 74))
            items = [
                "Log-distance path loss + log-normal shadowing：生成虚拟 RSSI/SNR",
                "固定 SF/BW/CR：当前结果控制 PHY 变量，便于看路由差异",
                "Half-duplex radio：正在发射的节点不能接收",
                "Collision + capture threshold：强信号可捕获，弱信号失败",
                "SNR margin -> probabilistic reception：链路不是硬阈值，而是概率成功",
            ]
            for j, item in enumerate(items):
                d.ellipse((190, 355 + j * 74, 208, 373 + j * 74), fill=(111, 186, 45) if j % 2 == 0 else (20, 104, 184))
                d.text((235, 338 + j * 74), item, font=pil_font(28), fill=(31, 41, 55))
            draw_center(d, (170, 825, 1580, 68), "来自文献的启发：ToA / 多 SF-aware 路由在异构拓扑中有机会提升 PDR、公平性与时延表现。", pil_font(29, True), (20, 104, 184))
        elif idx in (7, 8):
            source = unicast if idx == 7 else mixed
            msg = (
                f"Meshtastic-like PDR {source['meshtastic-like']['unicast_pdr']:.2f}，"
                f"MeshCore-like PDR {source['meshcore-like']['unicast_pdr']:.2f}\n"
                f"空口时间：{source['meshtastic-like']['total_airtime_s']:.0f}s → "
                f"{source['meshcore-like']['total_airtime_s']:.0f}s"
            )
            draw_center(d, (180, 335, 1560, 320), msg, pil_font(56, True), (17, 38, 74))
            d.text((270, 730), "关键解读：缓存路径减少冗余发送，但路径质量会影响单播可靠性。", font=pil_font(33, True), fill=(20, 104, 184))
        elif idx == 9:
            columns = [
                ("可靠性", ["Managed flooding 更稳", "重复副本增加成功概率", "代价是碰撞与冗余"], (20, 104, 184)),
                ("空口效率", ["Route cache 更省", "路径建立后只走必要节点", "代价是路径质量敏感"], (111, 186, 45)),
                ("适用流量", ["重复单播：缓存有价值", "广播/组播：泛洪仍重要", "高负载：需要拥塞感知"], (194, 48, 58)),
            ]
            for j, (head, items, color) in enumerate(columns):
                x = 180 + j * 560
                draw_center(d, (x, 250, 360, 55), head, pil_font(36, True), color)
                for k, item in enumerate(items):
                    d.ellipse((x + 10, 365 + k * 80, x + 28, 383 + k * 80), fill=color)
                    d.text((x + 50, 346 + k * 80), item, font=pil_font(28), fill=(31, 41, 55))
            draw_center(d, (170, 800, 1580, 85), "下一版路由应把 ToA、SNR margin、拥塞与路径新鲜度放进同一个成本函数。", pil_font(36, True), (17, 38, 74))
        elif idx == 10:
            boxes = [
                ("1. 加入链路代价", "把 ToA、SNR margin、重传风险、节点拥塞转换为路径成本。", (20, 104, 184)),
                ("2. 扩展 PHY 维度", "支持多 SF / 多信道，把准正交特性纳入路由与调度。", (111, 186, 45)),
                ("3. 做实证闭环", "用真实 SX1262 RSSI/SNR 日志校准传播模型。", (194, 48, 58)),
            ]
            for j, (head, body, color) in enumerate(boxes):
                x = 145 + j * 585
                d.rounded_rectangle((x, 245, x + 500, 650), radius=20, fill=(244, 249, 255), outline=(214, 228, 244), width=2)
                d.text((x + 34, 292), head, font=pil_font(32, True), fill=color)
                d.multiline_text((x + 34, 375), wrap_text(body, 19), font=pil_font(26), fill=(31, 41, 55), spacing=10)
            draw_center(d, (170, 775, 1580, 80), "目标：在不牺牲太多 PDR 的前提下，把空口占用和碰撞压下来。", pil_font(40, True), (17, 38, 74))
        else:
            d.text((170, 310), wrap_text("核心信息见 PPTX 中的可编辑图表、流程图和要点。", 26), font=pil_font(48, True), fill=(17, 38, 74))
        draw_footer(im, idx)
        slides.append(im)

    im = Image.open(assets["campus"]).convert("RGB").resize((1920, 1080))
    veil = Image.new("RGBA", im.size, (255, 255, 255, 160))
    im = Image.alpha_composite(im.convert("RGBA"), veil).convert("RGB")
    d = ImageDraw.Draw(im)
    draw_center(d, (0, 330, 1920, 170), "谢谢聆听", pil_font(82, True), (17, 38, 74))
    draw_center(d, (0, 510, 1920, 80), "欢迎交流 LoRa Mesh 路由、仿真参数与下一步实验设计", pil_font(34, True), (20, 104, 184))
    slides.append(im)

    for i, slide in enumerate(slides, 1):
        slide.save(PREVIEW_DIR / f"slide_{i:02d}.png")

    # Contact sheet for rhythm inspection.
    thumbs = []
    for slide in slides:
        t = slide.copy()
        t.thumbnail((320, 180))
        tile = Image.new("RGB", (340, 220), (245, 248, 252))
        tile.paste(t, (10, 10))
        thumbs.append(tile)
    contact = Image.new("RGB", (4 * 340, math.ceil(len(thumbs) / 4) * 220), (232, 238, 246))
    for i, tile in enumerate(thumbs):
        contact.paste(tile, ((i % 4) * 340, (i // 4) * 220))
    contact.save(PREVIEW_DIR / "contact_sheet.png")


def render_previews_v2(assets, unicast, mixed) -> None:
    for p in PREVIEW_DIR.glob("*.png"):
        p.unlink()

    preview_specs = [
        ("LoRa Mesh 路由仿真\n与协议取舍研究", "面向低速率、半双工无线网络的 PDR / 时延 / 空口效率比较"),
        ("先用一句话理解", "正在研究 LoRa 节点帮彼此转发消息时，怎样做到既送得到、又少占无线资源。重点看两点关系：转发方法和网络表现，可靠性和空口成本。"),
        ("核心结论先说人话", f"缓存路线能少发很多包：TX 下降 {pct_drop(unicast['meshtastic-like']['tx_count'], unicast['meshcore-like']['tx_count']):.0%}，空口下降 {pct_drop(unicast['meshtastic-like']['total_airtime_s'], unicast['meshcore-like']['total_airtime_s']):.0%}；但 PDR 从 {unicast['meshtastic-like']['unicast_pdr']:.2f} 降到 {unicast['meshcore-like']['unicast_pdr']:.2f}。"),
        ("目录", "00 开场 -> 01 核心结论 -> 02 背景与问题 -> 03 仿真内容 -> 04 两种转发思路 -> 05 仿真结果 -> 06 原因与改进"),
        ("为什么需要 LoRa Mesh", "LoRa 传得远，但只靠网关时覆盖和部署仍会受限制；mesh 可以接力转发，也会带来额外开销。"),
        ("主要看两件事", "大家都帮忙转发是不是更稳？先找好路线再发送能省多少资源？"),
        ("怎么比较才公平", "让两种方法在同样的节点、无线环境和流量里运行，再比较送达率、延迟和发包数量。"),
        ("仿真到底怎么跑", "输入节点、无线参数、流量和随机种子；过程按事件顺序计算发射、接收、碰撞和转发；输出 CSV 指标做平均比较。"),
        ("两种转发思路", "一种是“多几个人帮忙转”，更稳但更吵；一种是“先找路再沿路转”，更省但怕路线不好。"),
        ("仿真场景和参数", "50 nodes, 3000 m, 1200 s, 6 flows/min, pair-count 8, 20 seeds, SF9/125kHz/4/5。"),
        ("重复单播结果", f"MeshCore-like TX {unicast['meshcore-like']['tx_count']:.0f} vs {unicast['meshtastic-like']['tx_count']:.0f}；airtime {unicast['meshcore-like']['total_airtime_s']:.0f}s vs {unicast['meshtastic-like']['total_airtime_s']:.0f}s。"),
        ("混合流量结果", f"广播覆盖仍高：{mixed['meshcore-like']['broadcast_coverage']:.2f}；但 MeshCore-like 单播 PDR 降到 {mixed['meshcore-like']['unicast_pdr']:.2f}。"),
        ("原因分析", "缓存路线省掉了重复发送，也减少了从其他路径补救成功的机会。"),
        ("下一步怎么改", "不要只记住“有一条路”，还要判断路线短不短、稳不稳、忙不忙、新不新。"),
        ("后续工作", "扩展多 SF / 多信道，采集真实 SX1262 日志校准模型，补充置信区间与参数扫频。"),
        ("谢谢聆听", "欢迎交流 LoRa Mesh 路由、仿真参数与下一步实验设计。"),
    ]

    slides: list[Image.Image] = []
    for i, (title, body) in enumerate(preview_specs, 1):
        if i == 1:
            im = Image.open(assets["cover"]).convert("RGB").resize((1920, 1080))
            shade = Image.new("RGBA", im.size, (3, 19, 42, 92))
            im = Image.alpha_composite(im.convert("RGBA"), shade).convert("RGB")
            d = ImageDraw.Draw(im)
            d.multiline_text((120, 280), title, font=pil_font(76, True), fill=(255, 255, 255), spacing=16)
            d.text((124, 610), body, font=pil_font(30), fill=(224, 238, 255))
        elif i == len(preview_specs):
            im = Image.open(assets["campus"]).convert("RGB").resize((1920, 1080))
            veil = Image.new("RGBA", im.size, (255, 255, 255, 165))
            im = Image.alpha_composite(im.convert("RGBA"), veil).convert("RGB")
            d = ImageDraw.Draw(im)
            draw_center(d, (0, 330, 1920, 160), title, pil_font(82, True), (17, 38, 74))
            draw_center(d, (0, 520, 1920, 80), body, pil_font(34, True), (20, 104, 184))
        else:
            im = preview_base(assets)
            draw_title(im, title, "", f"{i-2:02d}")
            d = ImageDraw.Draw(im)
            d.multiline_text((170, 300), wrap_text(body, 34), font=pil_font(44, True), fill=(17, 38, 74), spacing=16)
            if i in (2, 11, 12):
                d.rounded_rectangle((220, 690, 1700, 810), radius=18, fill=(244, 249, 255), outline=(214, 228, 244), width=2)
                if i == 11:
                    msg = f"碰撞失败下降 {pct_drop(unicast['meshtastic-like']['collision_fail'], unicast['meshcore-like']['collision_fail']):.0%}"
                elif i == 12:
                    msg = f"混合场景空口下降 {pct_drop(mixed['meshtastic-like']['total_airtime_s'], mixed['meshcore-like']['total_airtime_s']):.0%}"
                else:
                    msg = "关键判断：缓存路径值得保留，但必须增加链路质量感知。"
                draw_center(d, (220, 700, 1480, 92), msg, pil_font(36, True), (20, 104, 184))
            draw_footer(im, i)
        slides.append(im)

    for i, slide in enumerate(slides, 1):
        slide.save(PREVIEW_DIR / f"slide_{i:02d}.png")

    thumbs = []
    for slide in slides:
        t = slide.copy()
        t.thumbnail((320, 180))
        tile = Image.new("RGB", (340, 220), (245, 248, 252))
        tile.paste(t, (10, 10))
        thumbs.append(tile)
    contact = Image.new("RGB", (4 * 340, math.ceil(len(thumbs) / 4) * 220), (232, 238, 246))
    for i, tile in enumerate(thumbs):
        contact.paste(tile, ((i % 4) * 340, (i // 4) * 220))
    contact.save(PREVIEW_DIR / "contact_sheet.png")


if __name__ == "__main__":
    build_pptx()
